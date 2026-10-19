# -*- coding: utf-8 -*-
"""產生圖解版操作說明（Word）。截圖在 shots/，由 a_*.py 用示範資料庫自動截出。"""
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SH = os.path.join(HERE, "shots")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "manual_illustrated.docx")
FONT = "Microsoft JhengHei"
RED = RGBColor(0xE5, 0x39, 0x35)
NAVY = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x59, 0x59, 0x59)

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = s.bottom_margin = Cm(1.8)


def font(style, size=None, bold=None, color=None):
    style.font.name = FONT
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = color


font(doc.styles["Normal"], 11.5)
doc.styles["Normal"].paragraph_format.space_after = Pt(4)
doc.styles["Normal"].paragraph_format.line_spacing = 1.25
font(doc.styles["Heading 1"], 20, True, NAVY)
font(doc.styles["Heading 2"], 14, True, NAVY)
for h in ("Heading 1", "Heading 2"):
    doc.styles[h].paragraph_format.space_before = Pt(10)
    doc.styles[h].paragraph_format.space_after = Pt(6)


def para(text="", bold=False, color=None, size=None, align=None, after=None):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold = bold or None
        if color:
            r.font.color.rgb = color
        if size:
            r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def img(name, width=17.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(SH, name), width=Cm(width))
    p.paragraph_format.space_after = Pt(6)


def shade(cell, fill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(sh)


def no_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def tiny():
    """表格後面必要的分隔段落：縮到極小，避免剛好滿版時被擠成一張空白頁"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = Pt(2)
    p.add_run("").font.size = Pt(1)


def legend(rows, widths=(1.3, 15.7)):
    """編號對照表：rows = [(編號, 名稱, 說明)]"""
    t = doc.add_table(rows=len(rows), cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, (n, name, desc) in enumerate(rows):
        c0, c1 = t.rows[i].cells
        c0.width, c1.width = Cm(widths[0]), Cm(widths[1])
        p = c0.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(n)
        r.bold = True
        r.font.color.rgb = RED
        r.font.size = Pt(13)
        p = c1.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(name)
        r.bold = True
        if desc:
            p.add_run("　" + desc)
        shade(c0, "FDECEA")
        no_split(t.rows[i])
    tiny()


def box(text, label, fill, label_color=None):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    c.width = Cm(17)
    shade(c, fill)
    p = c.paragraphs[0]
    r = p.add_run(label + "　")
    r.bold = True
    if label_color:
        r.font.color.rgb = label_color
    p.add_run(text)
    para(after=2)


def tip(text):
    box(text, "💡 小提醒", "FFF8E1")


def warn(text):
    box(text, "⚠️ 注意", "FDECEA", RED)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def h1(t, new_page=True):
    h = doc.add_heading(t, level=1)
    h.paragraph_format.page_break_before = new_page


def h2(t, new_page=False):
    h = doc.add_heading(t, level=2)
    h.paragraph_format.page_break_before = new_page


# ── 封面 ──
for _ in range(3):
    para()
para("股票紀錄系統", True, NAVY, 30, WD_ALIGN_PARAGRAPH.CENTER, 2)
para("操作說明（圖解版）", True, NAVY, 22, WD_ALIGN_PARAGRAPH.CENTER, 18)
para("看圖跟著做就會：每張圖上的紅色編號，對照下方說明即可", False, GRAY, 13, WD_ALIGN_PARAGRAPH.CENTER, 40)
para("內容", True, NAVY, 13, WD_ALIGN_PARAGRAPH.CENTER, 4)
for line in ["1. 登入", "2. 股票輸入（仿Yahoo）：買進、賣出、修改", "3. 當日交易明細", "4. 投資績效",
             "5. 已實現損益", "6. 庫存損益", "7. 個股明細", "8. 修改密碼", "9. 常用名詞與常見問題"]:
    para(line, False, None, 12, WD_ALIGN_PARAGRAPH.CENTER, 0)
para()
para("※ 文件中的畫面皆為示範資料，數字僅供說明。", False, GRAY, 10, WD_ALIGN_PARAGRAPH.CENTER)
page_break()

# ── 1. 登入 ──
h1("1. 登入", new_page=False)
img("p_login.png", 15)
legend([("①", "密碼", "只要輸入密碼，不用帳號。按右邊的眼睛可以看自己打了什麼"),
        ("②", "記住我", "預設已勾選：這台裝置 90 天內不用重新登入"),
        ("③", "登入", "密碼連續輸錯 5 次，要等 1 分鐘才能再試")])
tip("iPad／手機免登入：登入後進到「股票輸入（仿Yahoo）」頁，點 Safari 的「分享」→「加入主畫面」。"
    "之後從主畫面的圖示打開，90 天內直接進入，連密碼都不用打。")

# ── 2. 股票輸入 ──
h1("2. 股票輸入（仿Yahoo）")
para("最常用的一頁：記錄買賣、看持股。版面跟 Yahoo 股市的「持股明細」很像。")
img("y_overview.png", 12)
legend([("①", "分頁", "上方每一個分頁代表一個人；在哪個分頁輸入，交易就記在那個人名下"),
        ("②", "新增股票", "清單裡還沒有這檔股票時，按這裡打代號或名稱搜尋（例：3665 或 貿聯）；打錯一碼會提示「你是不是要找」"),
        ("③", "搜尋", "打股名或股號，馬上找到分頁裡的那一檔"),
        ("④", "輸入日期", "補登前幾天的交易時，先在這裡選日期（按鈕可一鍵選前兩個交易日），之後每一筆都自動帶這天；不是今天時會有黃色提醒"),
        ("⑤", "總覽", "持有股票市值、已實現損益、未實現損益（紅＝賺、綠＝賠）"),
        ("⑥", "更新股價", "盤中每 30 秒會自動更新；想立刻更新就按這裡"),
        ("⑦", "持股列", "點最左邊的 ▸ 展開，就可以在這檔下面記買賣"),
        ("⑧", "底價", "全部賣在這個價格，扣掉手續費和稅後剛好不虧")])

h2("2-1 記一筆「買進」", new_page=True)
img("y_buy.png", 17)
legend([("①", "日期", "預設用上方的「輸入日期」（平常就是今天），單筆也可以改"),
        ("②", "買／賣", "選「買入」"),
        ("③", "股數", "例如 1000"),
        ("④", "股價", "打完按 Enter"),
        ("⑤", "手續費、稅、成交金額", "系統自動算好，不用自己算"),
        ("⑥", "再加一筆買入", "同一檔一次要記好幾筆時使用"),
        ("⑦", "確認新增", "按鈕上會寫出買幾股、價格和日期，確認沒錯再按；送出後這檔會自動收合")])
tip("股價打完、游標還停在股價格子時，不要按鍵盤的 ↑ ↓：每按一下價格會加減 0.05 元。")

h2("2-2 記一筆「賣出」", new_page=True)
img("y_sell.png", 17)
legend([("①", "選「賣出」", ""),
        ("②", "當沖", "同一天買又賣，記得勾選（證交稅自動減半）"),
        ("③", "沖銷方式", "這筆賣出要對掉哪幾筆買進。預設「接近均價」，也可選先進先出、賺多、賠少等"),
        ("④", "時間範圍", "只看最近幾天買進的批次"),
        ("⑤", "沖銷表", "可直接修改最右邊「本次沖銷」的股數"),
        ("⑥", "檢查", "已配股數要等於賣出股數；這裡也會先算出這筆大約賺賠多少"),
        ("⑦", "確認新增", "")])

h2("2-3 改錯或刪除")
img("y_edit.png", 17)
legend([("①", "✏️ 修改", "可改日期、股數、股價；手續費和稅會自動重算"),
        ("②", "🗑 刪除", "會再問一次確認；要改「買／賣」方向請刪掉重新輸入")])

# ── 3~7 報表 ──
h1("3. 當日交易明細")
para("每天收盤後對帳、看交割金額用。")
img("r_daily.png", 11)
legend([("①", "分頁", "當日交易明細／區間交易明細／期間獲利表"),
        ("②", "篩選條件", "選日期、股票、買賣人"),
        ("③", "交割加總", "應收、應付、淨交割與當日已實現損益（黃底是重點）"),
        ("④", "每股當日損益簡表", "每檔一列，最下面深黃色是合計"),
        ("⑤", "每筆明細", "當天每一筆買賣"),
        ("⑥", "下載 CSV", "可用 Excel 打開")])

h1("4. 投資績效")
para("看一段期間的整體表現。")
img("r_perf.png", 12.5)
legend([("①", "快速區間", "當日、近3天、當週、當月、當年、近1年、全部"),
        ("②", "自訂日期", "也可以自己選開始、結束日期"),
        ("③", "顯示模式", "選要看已實現、未實現或合計"),
        ("④", "買賣人", "預設全部，可只看某一位"),
        ("⑤", "關鍵績效", "未實現、已實現、最佳個股、最差個股"),
        ("⑥", "其他數字", "總損益、勝率、平均每筆、盈虧比等")])
tip("每一個報表頁最上面的「篩選條件」用法都一樣，學會一頁就全部會了。")

h1("5. 已實現損益")
para("已經賣掉的部分賺賠多少，適合月底、年底結算。頁面下方還有每月趨勢圖、每一筆「買→賣」明細與匯出 Excel。")
img("r_realized.png", 10)
legend([("①", "快速區間", "例如選「當年」看今年全部"),
        ("②", "總已實現損益", "這段期間賺賠合計與報酬率"),
        ("③", "統計表", "勝率、盈虧比、手續費＋稅、當沖與波段損益"),
        ("④", "產業排名", "哪個產業賺最多，以及佔比"),
        ("⑤", "個股排名", "哪一檔賺最多，以及佔比")])

h1("6. 庫存損益")
para("目前手上還沒賣的股票。")
img("r_stock.png", 11.5)
legend([("①", "重點數字", "持倉市值、未實現損益、已實現損益、總報酬率"),
        ("②", "這些數字怎麼算？", "點開可以看公式和實際代入的數字"),
        ("③", "持倉明細", "每檔的股數、均價、底價（藍字）、現價、損益"),
        ("④", "底價怎麼算？", "點開看每一檔底價的計算過程")])

h1("7. 個股明細")
para("深入看單一檔股票。")
img("r_single.png", 10.5)
legend([("①", "選擇股票", "可直接打股號或股名搜尋"),
        ("②", "庫存", "還沒賣的買進明細與均價"),
        ("③", "已出售", "已賣出的部分：損益合計、獲利率、勝率與每一筆明細"),
        ("④", "匯出 Excel", "庫存＋已出售一次匯出")])

# ── 8. 修改密碼 ──
h1("8. 修改密碼")
warn("系統只有兩組密碼：「老闆」日常使用、「開發者」維護用。新密碼至少 8 碼，而且不能和另一組一樣（登入只輸入密碼，系統靠密碼分辨是誰）。"
     "改完後這台裝置要用新密碼重新登入；iPad 主畫面的圖示也要刪掉、重新「加入主畫面」。請先把新密碼寫下來。")
h2("步驟 1、2：打開「帳號與權限」")
t = doc.add_table(rows=1, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, (name, cap) in zip(t.rows[0].cells, [("p_side1.png", "① 選單最下面點「View more」展開"),
                                             ("p_side2.png", "② 點「帳號與權限」")]):
    c.width = Cm(8.5)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(SH, name), width=Cm(5.2))
    q = c.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.add_run(cap).bold = True
para()
h2("步驟 3、4：輸入新密碼並儲存")
img("p_card.png", 16)
legend([("✕", "帳號名稱", "這一格不要動"),
        ("①", "重設密碼", "點一下，輸入新密碼。打字時只會看到圓點，可按右邊的眼睛確認有沒有打錯"),
        ("②", "儲存帳號設定", "按下去")])
h2("步驟 5：看到綠色訊息就完成了")
img("p_done.png", 16)
legend([("③", "已更新", "出現這個訊息就代表成功")])
tip("建議馬上測試：點左下角「登出」，再用新密碼登入一次。萬一忘記密碼或登不進去，請找系統管理員重設，資料不會因此不見。")

# ── 9. 名詞與常見問題 ──
h1("9. 常用名詞與常見問題")
terms = [("均價", "手上這檔股票平均每股成本（含買進手續費）"),
         ("底價（藍字）", "全部賣在這個價，扣掉手續費和證交稅後剛好不虧；會比均價高一點"),
         ("沖銷", "把一筆賣出和之前的買進對起來，才算得出賺賠"),
         ("已實現損益", "已經賣掉的部分實際賺賠（已扣手續費、稅）"),
         ("未實現損益", "還沒賣的股票，用現在股價算的賺賠"),
         ("交割應收／應付", "賣出拿回／買進要付的錢（含手續費、稅）"),
         ("當沖", "同一天買進又賣出，證交稅減半（0.15%）"),
         ("手續費", "成交金額 × 0.1425% × 折數，小數捨去，最低 1 元"),
         ("證交稅", "賣出才收：一般 0.3%、當沖 0.15%、ETF 0.1%"),
         ("紅色／綠色", "台股習慣：紅色＝賺、綠色＝賠")]
t = doc.add_table(rows=len(terms) + 1, cols=2)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (a, b_) in enumerate([("名詞", "意思")] + terms):
    c0, c1 = t.rows[i].cells
    c0.width, c1.width = Cm(4), Cm(13)
    c0.paragraphs[0].add_run(a).bold = True
    c1.paragraphs[0].add_run(b_).bold = (i == 0) or None
    if i == 0:
        shade(c0, "D9E2F3")
        shade(c1, "D9E2F3")
para()
h2("遇到問題時")
for q, a in [("畫面出現紅色錯誤訊息", "先按瀏覽器的「重新整理」；還是一樣就截圖給系統管理員。"),
             ("金額跟券商差一點點", "先逐筆對照成交價（最常見是某一筆差 0.05 元），再對照手續費。"),
             ("股價沒有更新", "收盤後股價不會變；盤中可按「更新股價」。"),
             ("每次打開都要重新登入", "iPad 請在「股票輸入（仿Yahoo）」頁用 Safari「分享 → 加入主畫面」，之後從主畫面圖示打開；改過密碼要重新加一次。"),
             ("打開不是在「股票輸入」頁", "主畫面圖示或書籤存到別的頁面了，請在「股票輸入（仿Yahoo）」頁重新加入。")]:
    p = doc.add_paragraph()
    r = p.add_run("● " + q + "：")
    r.bold = True
    r.font.color.rgb = NAVY
    p.add_run(a)
    p.paragraph_format.space_after = Pt(3)

# 頁尾頁碼
fp = doc.sections[0].footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    e = OxmlElement("w:fldChar" if tag else "w:instrText")
    if tag:
        e.set(qn("w:fldCharType"), tag)
    else:
        e.text = txt
    run._r.append(e)
doc.save(OUT)
print("saved", OUT)
