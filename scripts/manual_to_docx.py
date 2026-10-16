# -*- coding: utf-8 -*-
"""把 docs/操作手冊.md 轉成 Word 檔：python scripts/manual_to_docx.py docs/操作手冊.md 輸出.docx"""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, OUT = sys.argv[1], sys.argv[2]
FONT = "Microsoft JhengHei"
doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2)

def set_font(style, size=None, bold=None, color=None):
    style.font.name = FONT
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size: style.font.size = Pt(size)
    if bold is not None: style.font.bold = bold
    if color: style.font.color.rgb = RGBColor.from_string(color)
set_font(doc.styles["Normal"], 11)
doc.styles["Normal"].paragraph_format.space_after = Pt(4)
doc.styles["Normal"].paragraph_format.line_spacing = 1.3
for name, size in (("Title", 24), ("Heading 1", 17), ("Heading 2", 14), ("Heading 3", 12)):
    set_font(doc.styles[name], size, True, "1F3864")
for name in ("List Bullet", "List Number"):
    set_font(doc.styles[name], 11)

def add_runs(p, text, base_bold=False):
    # **粗體**、`程式字`
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part: continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"
        else:
            r = p.add_run(part); r.bold = base_bold or None

def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_)
    tcPr.append(sh)

def hr(p):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "BFBFBF")): bt.set(qn(k), v)
    b.append(bt); pPr.append(b)

lines = open(SRC, encoding="utf-8").read().splitlines()
i = 0
num_restart = True
while i < len(lines):
    ln = lines[i].rstrip()
    if not ln.strip():
        num_restart = True; i += 1; continue
    if ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells): rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = t.cell(r, c); cell.text = ""
                p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
                add_runs(p, val, base_bold=(r == 0))
                for run in p.runs: run.font.size = Pt(10)
                if r == 0: shade(cell, "D9E2F3")
        # 欄寬：中間說明欄最寬（總寬 16.6cm）
        n = len(rows[0])
        widths = {2: [5.5, 11.1], 3: [3.6, 9.2, 3.8]}.get(n, [16.6 / n] * n)
        t.autofit = False
        for r_ in t.rows:
            trPr = r_._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
            for c, w in zip(r_.cells, widths):
                c.width = Cm(w)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        continue
    m = re.match(r"^(#{1,3}) (.*)", ln)
    if m:
        lvl, txt = len(m.group(1)), m.group(2).replace("📘 ", "")
        if lvl == 1:
            h = doc.add_paragraph(style="Title"); add_runs(h, txt)
        else:
            h = doc.add_heading(level=lvl - 1); add_runs(h, txt)
        i += 1; continue
    if ln.strip() == "---":
        hr(doc.add_paragraph()); i += 1; continue
    if ln.startswith(">"):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6)
        add_runs(p, ln.lstrip("> ").strip())
        for r in p.runs: r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        i += 1; continue
    m = re.match(r"^(\s*)- (.*)", ln)
    if m:
        p = doc.add_paragraph(style="List Bullet")
        if m.group(1): p.paragraph_format.left_indent = Cm(1.6)
        add_runs(p, m.group(2)); i += 1; continue
    m = re.match(r"^(\d+)\. (.*)", ln)
    if m:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.6)
        add_runs(p, f"{m.group(1)}.  " + m.group(2)); i += 1; continue
    p = doc.add_paragraph(); add_runs(p, ln.strip()); i += 1

# 頁尾頁碼
sec = doc.sections[0]; fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run()
for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    if tag:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r._r.append(e)
    else:
        e = OxmlElement("w:instrText"); e.text = txt; r._r.append(e)
doc.save(OUT)
print("saved", OUT)
